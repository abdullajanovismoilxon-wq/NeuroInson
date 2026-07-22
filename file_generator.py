# file_generator.py — Automatic PPTX, DOCX, ZIP generator for Neuroinson Project Center
import os
import re
import json
import time
import zipfile
import py_compile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

GENERATED_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generated_files")
os.makedirs(GENERATED_DIR, exist_ok=True)

def cleanup_old_files(max_age_seconds=86400):
    """Deletes files older than 24 hours."""
    now = time.time()
    try:
        for fname in os.listdir(GENERATED_DIR):
            fpath = os.path.join(GENERATED_DIR, fname)
            if os.path.isfile(fpath) and (now - os.path.getmtime(fpath) > max_age_seconds):
                os.remove(fpath)
    except Exception as e:
        print("[CLEANUP ERR]", e)

def generate_pptx_file(text, user_id=0):
    """Parses JSON or markdown from text and creates a real PowerPoint .pptx presentation."""
    cleanup_old_files()
    filename = f"taqdimot_{user_id}_{int(time.time())}.pptx"
    filepath = os.path.join(GENERATED_DIR, filename)

    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Try JSON extraction
    json_data = None
    json_match = re.search(r'```(?:json)?\s*(\{[\s\S]*?\})\s*```', text)
    if json_match:
        try:
            json_data = json.loads(json_match.group(1))
        except Exception:
            pass
    
    if not json_data:
        try:
            json_data = json.loads(text)
        except Exception:
            pass

    if not json_data:
        # Fallback markdown parser into structured slides
        slides_list = []
        raw_slides = re.split(r'(?i)(?:slayd\s*\d+:?|\#\#?\s*Slayd\s*\d+:?)', text)
        title = "Taqdimot"
        for idx, block in enumerate(raw_slides):
            block = block.strip()
            if not block:
                continue
            lines = [l.strip() for l in block.split('\n') if l.strip()]
            if idx == 0 and len(lines) > 0:
                title = lines[0].replace('#', '').strip()
                continue
            stitle = lines[0].replace('*', '').replace('#', '').strip()
            bullets = [l.replace('-', '').replace('*', '').strip() for l in lines[1:] if l.startswith('-') or l.startswith('*')]
            if not bullets:
                bullets = lines[1:5]
            slides_list.append({"title": stitle, "bullets": bullets, "notes": ""})
        json_data = {"title": title, "subtitle": "Neuroinson AI tomonidan tayyorlandi", "slides": slides_list}

    # Add Title Slide
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Title Slide Background
    bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5)) # 1 = MSO_SHAPE.RECTANGLE
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42) # Deep slate/dark blue
    bg.line.fill.background()

    # Title Text Box
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.333), Inches(3.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = json_data.get("title", "Taqdimot Slaydi")
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = json_data.get("subtitle", "Neuroinson AI Presentation Center")
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(99, 102, 241) # Indigo accent
    p2.font.name = "Arial"
    p2.alignment = PP_ALIGN.CENTER

    # Content Slides
    slides = json_data.get("slides", [])
    for slide_info in slides:
        c_slide = prs.slides.add_slide(blank_layout)
        # Background
        c_bg = c_slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        c_bg.fill.solid()
        c_bg.fill.fore_color.rgb = RGBColor(248, 250, 252) # Clean light theme for content
        c_bg.line.fill.background()

        # Header Box
        header_box = c_slide.shapes.add_shape(1, Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.2))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
        header_box.line.fill.background()
        
        htf = header_box.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = slide_info.get("title", "Slayd")
        hp.font.size = Pt(26)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(255, 255, 255)
        hp.font.name = "Arial"
        hp.alignment = PP_ALIGN.LEFT

        # Bullets Box
        b_box = c_slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.8))
        btf = b_box.text_frame
        btf.word_wrap = True

        bullets = slide_info.get("bullets", [])
        for b_idx, b_text in enumerate(bullets):
            bp = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
            bp.text = f"•  {b_text}"
            bp.font.size = Pt(20)
            bp.font.color.rgb = RGBColor(51, 65, 85)
            bp.font.name = "Arial"
            bp.space_after = Pt(14)

        # Speaker notes if present
        notes_text = slide_info.get("notes", "")
        if notes_text:
            notes_slide = c_slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = notes_text

    prs.save(filepath)
    return f"/generated_files/{filename}"

def generate_docx_file(text, user_id=0):
    """Creates a real Microsoft Word .docx document for Kurs Ishi / Mustaqil Ish."""
    cleanup_old_files()
    filename = f"kurs_ishi_{user_id}_{int(time.time())}.docx"
    filepath = os.path.join(GENERATED_DIR, filename)

    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = DocxInches(1)
        section.bottom_margin = DocxInches(1)
        section.left_margin = DocxInches(1.2)
        section.right_margin = DocxInches(0.8)

    # Base style setup
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = DocxPt(14)

    # Extract title or main heading
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    doc_title = "KURS ISHI"
    if lines:
        doc_title = lines[0].replace('#', '').strip()

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(doc_title.upper())
    run.bold = True
    run.font.size = DocxPt(18)
    run.font.color.rgb = DocxRGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = DocxPt(24)

    # Parse and write content lines
    in_code_block = False
    for line in lines[1:]:
        if line.startswith('```'):
            in_code_block = not in_code_block
            continue
        
        if line.startswith('#'):
            level = line.count('#')
            htext = line.replace('#', '').strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = DocxPt(16)
            hp.paragraph_format.space_after = DocxPt(8)
            hrun = hp.add_run(htext)
            hrun.bold = True
            hrun.font.size = DocxPt(16 if level == 1 else 14)
            hrun.font.color.rgb = DocxRGBColor(30, 41, 59)
        elif line.startswith('- ') or line.startswith('* '):
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = DocxPt(4)
            brun = bp.add_run(line[2:].strip())
            brun.font.size = DocxPt(14)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = DocxPt(8)
            prun = p.add_run(line.strip())
            prun.font.size = DocxPt(14)

    doc.save(filepath)
    return f"/generated_files/{filename}"

def generate_webapp_zip(text, user_id=0):
    """Parses HTML/CSS/JS from text and creates a downloadable .zip archive for Web App."""
    cleanup_old_files()
    filename = f"web_app_{user_id}_{int(time.time())}.zip"
    filepath = os.path.join(GENERATED_DIR, filename)

    # Extract code blocks
    html_code = ""
    css_code = ""
    js_code = ""

    html_match = re.search(r'```html\n([\s\S]*?)```', text, re.IGNORECASE)
    if html_match:
        html_code = html_match.group(1).strip()
    else:
        # Fallback if unlabelled HTML block
        fallback_match = re.search(r'```\n([\s\S]*?<!DOCTYPE html[\s\S]*?)```', text, re.IGNORECASE)
        if fallback_match:
            html_code = fallback_match.group(1).strip()

    css_match = re.search(r'```css\n([\s\S]*?)```', text, re.IGNORECASE)
    if css_match:
        css_code = css_match.group(1).strip()

    js_match = re.search(r'```javascript\n([\s\S]*?)```', text, re.IGNORECASE) or re.search(r'```js\n([\s\S]*?)```', text, re.IGNORECASE)
    if js_match:
        js_code = js_match.group(1).strip()

    if not html_code:
        html_code = f"<!DOCTYPE html>\n<html lang=\"uz\">\n<head>\n<meta charset=\"UTF-8\">\n<title>Web App</title>\n" \
                    f"<style>\n{css_code or 'body { font-family: sans-serif; padding: 20px; }'}\n</style>\n</head>\n<body>\n" \
                    f"<pre>{text}</pre>\n<script>\n{js_code}\n</script>\n</body>\n</html>"

    readme_content = ("=========================================\n"
                      "  NEUROINSON AI — WEB APP LOYIHASI\n"
                      "=========================================\n\n"
                      "O'RNATISH VA ISHGATUSHIRISH QO'LLANMASI:\n\n"
                      "1. Ushbu ZIP arxivini ixtiyoriy papkaga chiqaring (Extract).\n"
                      "2. 'index.html' faylini kompyuteringizdagi istalgan brauzerda (Chrome, Edge, Firefox) oching.\n"
                      "3. Loyiha avtomatik ishga tushadi!\n\n"
                      "Muvaffaqiyatli foydalanishni tilaymiz! (Neuroinson AI Hub)\n")

    with zipfile.ZipFile(filepath, 'w', zipfile.ZIP_DEFLATED) as zipf:
        zipf.writestr("index.html", html_code)
        if css_code:
            zipf.writestr("style.css", css_code)
        if js_code:
            zipf.writestr("script.js", js_code)
        zipf.writestr("README.txt", readme_content)

    return f"/generated_files/{filename}"

def generate_bot_zip(text, user_id=0):
    """Parses Python code from text, verifies syntax, and creates a downloadable .zip archive for Telegram Bot."""
    cleanup_old_files()
    filename = f"telegram_bot_{user_id}_{int(time.time())}.zip"
    filepath = os.path.join(GENERATED_DIR, filename)

    python_code = ""
    py_match = re.search(r'```python\n([\s\S]*?)```', text, re.IGNORECASE) or re.search(r'```py\n([\s\S]*?)```', text, re.IGNORECASE)
    if py_match:
        python_code = py_match.group(1).strip()
    else:
        # Fallback extract code block
        cb_match = re.search(r'```\n([\s\S]*?)```', text)
        if cb_match:
            python_code = cb_match.group(1).strip()

    if not python_code:
        python_code = "# Telegram Bot Code\n" + text

    # Syntax compile check
    syntax_ok = True
    try:
        compile(python_code, '<string>', 'exec')
    except Exception as e:
        syntax_ok = False
        python_code += f"\n\n# Note: Syntax check warning: {str(e)}\n"

    req_content = ("pyTelegramBotAPI>=4.14.0\n"
                   "requests>=2.31.0\n"
                   "python-dotenv>=1.0.0\n")

    readme_content = ("# Telegram Bot Loyihasi (Neuroinson AI)\n\n"
                      "## O'rnatish va Ishga tushirish:\n\n"
                      "1. Kompyuterda Python 3.9+ o'rnatilganini tekshiring.\n"
                      "2. Terminal/CMD'da quyidagi buyruq orqali kutubxonalarni o'rnating:\n"
                      "   `pip install -r requirements.txt`\n"
                      "3. `bot.py` faylidagi BOT_TOKEN o'zgaruvchisiga @BotFather'dan olingan tokenni kiriting.\n"
                      "4. Botni ishga tushiring:\n"
                      "   `python bot.py`\n")

    with zipfile.ZipFile(filepath, 'w', zipfile.ZIP_DEFLATED) as zipf:
        zipf.writestr("bot.py", python_code)
        zipf.writestr("requirements.txt", req_content)
        zipf.writestr("README.md", readme_content)

    return f"/generated_files/{filename}"

if __name__ == '__main__':
    import sys
    if len(sys.argv) > 2 and sys.argv[1] == '--json':
        input_json_file = sys.argv[2]
        if os.path.exists(input_json_file):
            try:
                with open(input_json_file, 'r', encoding='utf-8') as f:
                    pdata = json.load(f)
                reply_text = pdata.get('reply', '')
                user_id = pdata.get('user_id', 0)
                proj_name = pdata.get('project_name', '').lower()

                f_url = None
                f_label = None

                if 'taqdimot' in proj_name or 'pptx' in proj_name:
                    f_url = generate_pptx_file(reply_text, user_id)
                    f_label = "📥 Taqdimot Faylini Yuklab Olish (.pptx)"
                elif 'kurs' in proj_name or 'mustaqil' in proj_name:
                    f_url = generate_docx_file(reply_text, user_id)
                    f_label = "📥 Kurs Ishi Hujjatini Yuklab Olish (.docx)"
                elif 'web app' in proj_name or 'sayt' in proj_name:
                    f_url = generate_webapp_zip(reply_text, user_id)
                    f_label = "📥 Sayt Kodini Yuklab Olish (.zip)"
                elif 'bot' in proj_name:
                    f_url = generate_bot_zip(reply_text, user_id)
                    f_label = "📥 Telegram Bot Loyihasini Yuklab Olish (.zip)"

                print(json.dumps({"file_url": f_url, "file_label": f_label}))
            except Exception as e:
                print(json.dumps({"error": str(e)}))

